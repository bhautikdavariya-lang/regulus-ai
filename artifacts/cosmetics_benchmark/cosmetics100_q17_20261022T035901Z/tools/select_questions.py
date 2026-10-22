"""Select 100 questions from Regulus_QBank_Cosmetics_REV01.xlsx (Questions sheet), reproducibly.

    PYTHONPATH=<openpyxl dir> python select_questions.py <bank.xlsx> <out_dir> [--seed 20261020] [--n 100]

Method: stratified random sampling by the bank's own Part classification (column B). Rows whose Part
is blank (the 24 UC-COS cross-cutting questions) form their own stratum; nothing is re-classified.
Allocation is proportional with the largest-remainder rule; ties on the remainder are broken by
larger stratum, then by stratum label. Within a stratum, questions are drawn with
random.Random(f"{seed}:{stratum}").sample, so each stratum's draw is independent of the others.
Question text, ID and key columns are copied verbatim. The bank is opened read-only and its sha256
is recorded before and after.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

KEEP = ["ID", "Part", "Subpart", "Section", "Section Heading", "Focus", "Question Type", "Difficulty",
        "Question (ask verbatim)", "Expected Citation — VERIFIED KEY", "Key Answer Elements",
        "Key Verified (Y/N)", "Verification Evidence (sections read; current eCFR)", "Audit Correction Note"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("bank")
    ap.add_argument("out_dir")
    ap.add_argument("--seed", type=int, default=20261020)
    ap.add_argument("--n", type=int, default=100)
    a = ap.parse_args()
    bank, out = Path(a.bank), Path(a.out_dir)
    sha_before = sha256(bank)

    wb = load_workbook(bank, read_only=True)
    rows = list(wb["Questions"].iter_rows(values_only=True))
    wb.close()
    hdr = list(rows[0])
    data = [dict(zip(hdr, r)) for r in rows[1:] if r and r[0]]

    # eligibility: non-blank ID and question, unique ID, unique normalised question text
    ids = Counter(d["ID"] for d in data)
    norm = Counter(" ".join(str(d["Question (ask verbatim)"] or "").lower().split()) for d in data)
    excluded = []
    eligible = []
    for d in data:
        q = d["Question (ask verbatim)"]
        n = " ".join(str(q or "").lower().split())
        if not q or not str(q).strip():
            excluded.append({"id": d["ID"], "reason": "blank question"})
        elif ids[d["ID"]] > 1:
            excluded.append({"id": d["ID"], "reason": "duplicate ID"})
        elif norm[n] > 1:
            excluded.append({"id": d["ID"], "reason": "duplicate question text"})
        else:
            eligible.append(d)
    if len(eligible) < a.n:
        print(f"STOP: only {len(eligible)} eligible unique questions (< {a.n})")
        return 2

    def stratum(d):
        return str(d["Part"]) if d["Part"] not in (None, "") else "UC-COS (Part blank)"

    groups: dict[str, list] = {}
    for d in eligible:
        groups.setdefault(stratum(d), []).append(d)
    total = len(eligible)
    quota = {s: a.n * len(g) / total for s, g in groups.items()}
    alloc = {s: int(q) for s, q in quota.items()}
    order = sorted(groups, key=lambda s: (-(quota[s] - alloc[s]), -len(groups[s]), s))
    for s in order[: a.n - sum(alloc.values())]:
        alloc[s] += 1

    selected = []
    for s in sorted(groups):
        pool = groups[s]   # bank order
        pick = random.Random(f"{a.seed}:{s}").sample(range(len(pool)), alloc[s])
        for i in sorted(pick):
            selected.append({**{k: pool[i].get(k) for k in KEEP}, "Stratum": s})
    # run order = bank order, so the run is reproducible and not grouped by stratum choice
    bank_pos = {d["ID"]: i for i, d in enumerate(data)}
    selected.sort(key=lambda r: bank_pos[r["ID"]])
    for i, r in enumerate(selected, start=1):
        r["Run Order"] = i

    sha_after = sha256(bank)
    out.mkdir(parents=True, exist_ok=True)
    summary = {
        "bank_file": bank.name, "bank_sha256_before": sha_before, "bank_sha256_after": sha_after,
        "bank_unchanged": sha_before == sha_after, "sheet": "Questions",
        "bank_rows": len(data), "eligible": len(eligible), "excluded": excluded,
        "method": "stratified random sampling by Part (column B; blank Part = own stratum), proportional "
                  "largest-remainder allocation (ties: larger stratum, then label), "
                  "random.Random(f'{seed}:{stratum}').sample within stratum; run in bank order",
        "seed": a.seed, "n": a.n,
        "allocation": {s: {"bank": len(groups[s]), "quota": round(quota[s], 3), "selected": alloc[s]} for s in sorted(groups)},
        "distribution": {
            col: {"bank": dict(Counter(str(d[col]) for d in eligible)),
                  "selected": dict(Counter(str(r[col]) for r in selected))}
            for col in ("Difficulty", "Question Type")},
        "selected_ids": [r["ID"] for r in selected],
        "python": sys.version.split()[0],
    }
    (out / "selection.json").write_text(json.dumps({"summary": summary, "questions": selected}, indent=1,
                                                   ensure_ascii=False, default=str), encoding="utf-8")
    with (out / "selected_questions.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["Run Order", "Stratum"] + KEEP)
        w.writeheader()
        w.writerows(selected)
    print(json.dumps({k: v for k, v in summary.items() if k != "selected_ids"}, indent=1, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
