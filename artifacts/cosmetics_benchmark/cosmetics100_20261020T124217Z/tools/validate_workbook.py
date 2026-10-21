"""Reopen the workbook, check integrity against the raw run, and evaluate every Summary formula with pycel."""
import json, sys
from pathlib import Path
from openpyxl import load_workbook
from pycel import ExcelCompiler

run, xlsx = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
exp = json.loads((xlsx.parent / (xlsx.stem + "_expected_values.json")).read_text())
results = {json.loads(l)["question_id"]: json.loads(l) for l in (run / "results.jsonl").read_text().splitlines() if l.strip()}
sel = json.loads((run / "selection.json").read_text())
bank = {q["ID"]: q for q in sel["questions"]}
problems = []
wb = load_workbook(xlsx)
need = ["Question Results", "Summary", "Failed and Partial Questions", "Run Metadata"]
if wb.sheetnames != need:
    problems.append(f"sheets {wb.sheetnames}")
ws = wb["Question Results"]
hdr = [c.value for c in ws[1]]
req = ["Question ID", "Question", "Reference Answer", "Generated Answer", "Evaluation Status", "Evaluation Explanation",
       "Duration (seconds)", "Input Tokens", "Output Tokens", "Base Model", "Embedding Model", "Execution Status",
       "Error Details", "Retrieval Notes"]
if hdr[:14] != req:
    problems.append(f"header {hdr[:14]}")
col = {h: i for i, h in enumerate(hdr)}
rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
ids = [r[0] for r in rows]
if len(rows) != 100 or len(set(ids)) != 100 or set(ids) != set(sel["summary"]["selected_ids"]):
    problems.append("row/ID mismatch")
for r in rows:
    q, res = bank[r[0]], results[r[0]]
    if r[col["Question"]] != q["Question (ask verbatim)"]: problems.append(f"{r[0]} question text differs")
    if r[col["Reference Answer"]] != q["Key Answer Elements"]: problems.append(f"{r[0]} reference differs")
    if r[col["Generated Answer"]] != res["answer"]: problems.append(f"{r[0]} answer differs/truncated")
    if r[col["Generated Answer"]] == q["Key Answer Elements"]: problems.append(f"{r[0]} answer equals reference")
    if not isinstance(r[col["Duration (seconds)"]], float): problems.append(f"{r[0]} duration not numeric")
    for t in ("Input Tokens", "Output Tokens"):
        if not isinstance(r[col[t]], int): problems.append(f"{r[0]} {t} not int")
    if r[col["Evaluation Status"]] not in ("PASS", "PARTIALLY PASS", "PARTIALLY FAIL", "FAIL", "NOT EVALUATED"): problems.append(f"{r[0]} bad status")
    if not r[col["Evaluation Explanation"]]: problems.append(f"{r[0]} no explanation")
    if len(str(r[col["Generated Answer"]] or "")) > 32767: problems.append(f"{r[0]} exceeds cell limit")
fp = wb["Failed and Partial Questions"]
fp_ids = [r[0] for r in fp.iter_rows(min_row=2, values_only=True) if r[0]]
non_pass = [r[0] for r in rows if r[col["Evaluation Status"]] != "PASS"]
if sorted(fp_ids) != sorted(non_pass): problems.append("failed/partial sheet does not match non-PASS rows")
if any(not r[9] for r in fp.iter_rows(min_row=2, values_only=True) if r[0]): problems.append("missing review action")

# evaluate every Summary formula
xc = ExcelCompiler(filename=str(xlsx), plugins=["xl_extra"])
sm = wb["Summary"]
values = {}
errors = []
for row in sm.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            v = xc.evaluate(f"Summary!{c.coordinate}")
            label = sm.cell(row=c.row, column=2).value
            values[(label, c.column_letter)] = v
            if isinstance(v, str) and v.startswith("#") or v is None:
                errors.append((c.coordinate, c.value, v))
def get(label, colL="C"):
    return values[(label, colL)]
checks = {
 "Total questions selected": 100, "Total questions attempted": 100,
 "Total successfully generated answers": exp["success"], "  of which withheld by the validator": exp["withheld"],
 "Total execution errors": 0, "Total timeouts": 0, "Total evaluated questions": exp["evaluated"], "Questions not evaluated": exp["counts"]["NOT EVALUATED"],
 "PASS": exp["counts"]["PASS"], "PARTIALLY PASS": exp["counts"]["PARTIALLY PASS"], "PARTIALLY FAIL": exp["counts"]["PARTIALLY FAIL"],
 "FAIL": exp["counts"]["FAIL"], "NOT EVALUATED": exp["counts"]["NOT EVALUATED"], "TOTAL SELECTED": 100,
 "Strict pass rate": exp["counts"]["PASS"] / exp["evaluated"],
 "Pass-or-partially-pass rate": (exp["counts"]["PASS"] + exp["counts"]["PARTIALLY PASS"]) / exp["evaluated"],
 "Partial-failure rate": exp["counts"]["PARTIALLY FAIL"] / exp["evaluated"], "Failure rate": exp["counts"]["FAIL"] / exp["evaluated"],
 "Evaluation coverage": exp["evaluated"] / 100, "Execution success rate": exp["success"] / 100,
 "Total benchmark duration — sum of question durations (s)": exp["sum_dur"],
 "Average question duration (s)": exp["avg_dur"], "Median question duration (s)": exp["median_dur"],
 "Minimum question duration (s)": exp["min_dur"], "Maximum question duration (s)": exp["max_dur"], "Duration observations": exp["n_dur"],
 "Average input tokens": exp["avg_in"], "Average output tokens": exp["avg_out"], "Total input tokens": exp["sum_in"],
 "Total output tokens": exp["sum_out"], "Input-token observations": exp["n_in"], "Output-token observations": exp["n_out"],
 "Questions missing duration": 0, "Questions missing input tokens": 0, "Questions missing output tokens": 0,
 "Live evaluator: PASS": exp["live"]["PASS"], "Live evaluator: FAIL": exp["live"]["FAIL"],
 "Live evaluator: NOT EVALUATED": exp["live"]["NOT EVALUATED"], "Live evaluator: NOT FULLY EVALUATED": exp["live"]["NOT FULLY EVALUATED"],
}
for label, want in checks.items():
    got = get(label)
    if abs(float(got) - float(want)) > 1e-9:
        problems.append(f"Summary '{label}': formula {got} != expected {want}")
recon = values[("TOTAL SELECTED", "D")]
if recon != "reconciles with questions selected": problems.append(f"reconciliation: {recon}")
print(json.dumps({"formulas_evaluated": len(values), "formula_errors": errors, "problems": problems,
                  "sample": {k: get(k) for k in ("Strict pass rate", "Pass-or-partially-pass rate", "Average question duration (s)",
                                                  "Median question duration (s)", "Total input tokens", "Total output tokens")}},
                 indent=1, default=str))
sys.exit(1 if problems or errors else 0)
